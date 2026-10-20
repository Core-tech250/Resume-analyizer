import io
import logging
from typing import Tuple
from docx import Document

logger = logging.getLogger(__name__)

def parse_docx(file_bytes: bytes) -> Tuple[str, int]:
    """
    Extracts text from DOCX bytes including paragraphs, headings, and tables.
    Returns (extracted_text, paragraph_count).
    """
    try:
        doc = Document(io.BytesIO(file_bytes))
        sections = []

        # Extract text from paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                sections.append(text)

        # Extract text from tables (common in resumes)
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    sections.append(" | ".join(row_texts))

        full_text = "\n".join(sections).strip()
        return full_text, len(doc.paragraphs)
    except Exception as e:
        logger.error(f"DOCX extraction failed: {e}")
        raise ValueError(f"Could not parse DOCX file: {e}")
